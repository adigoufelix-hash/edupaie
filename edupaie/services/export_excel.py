from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from edupaie.services.recu_pdf import MODES

ENTETE = PatternFill("solid", fgColor="1E1B4B")
TOTAL = PatternFill("solid", fgColor="EEF2FF")


def _chemin(dossier, prefixe) -> Path:
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    return dossier / f"{prefixe}_{datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx"


def _date(valeur):
    return date.fromisoformat(valeur[:10]) if valeur else None


def _cellule(ws, ligne, colonne, valeur, genre):
    if genre == "texte":
        valeur = "" if valeur is None else str(valeur)
    c = ws.cell(row=ligne, column=colonne, value=valeur)
    if genre == "texte":
        c.data_type = "s"  # jamais interprété comme une formule
    elif genre == "montant":
        c.number_format = "#,##0"
        c.alignment = Alignment(horizontal="right")
    elif genre == "date":
        c.number_format = "DD/MM/YYYY"
        c.alignment = Alignment(horizontal="center")
    elif genre == "nombre":
        c.alignment = Alignment(horizontal="center")
    return c


def _ecrire(ws, colonnes, lignes, total=None):
    """colonnes : [(titre, genre)] ; lignes : listes de valeurs ; total : liste ou None."""
    for j, (titre, _) in enumerate(colonnes, start=1):
        c = ws.cell(row=1, column=j, value=titre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENTETE
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, valeurs in enumerate(lignes, start=2):
        for j, ((_, genre), v) in enumerate(zip(colonnes, valeurs), start=1):
            _cellule(ws, i, j, v, genre)
    if total:
        for j, ((_, genre), v) in enumerate(zip(colonnes, total), start=1):
            c = _cellule(ws, len(lignes) + 2, j, v, genre)
            c.font = Font(bold=True)
            c.fill = TOTAL
    for j, (titre, _) in enumerate(colonnes, start=1):
        largeur = len(titre)
        for rangee in ws.iter_rows(min_row=2, min_col=j, max_col=j):
            v = rangee[0].value
            largeur = max(largeur, len(f"{v:,}") if isinstance(v, int) else len(str(v or "")))
        ws.column_dimensions[get_column_letter(j)].width = min(largeur + 3, 45)
    ws.freeze_panes = "A2"
    if lignes:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colonnes))}{len(lignes) + 1}"


def exporter_eleves(situations, dossier) -> Path:
    """situations : liste de (élève, situation), comme PaiementService.situations()."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Élèves"
    colonnes = [("Nom", "texte"), ("Prénom", "texte"), ("Classe", "texte"),
                ("Année scolaire", "texte"), ("Frais (FCFA)", "montant"),
                ("Remises (FCFA)", "montant"), ("Total dû (FCFA)", "montant"),
                ("Payé (FCFA)", "montant"), ("Solde (FCFA)", "montant"),
                ("Statut", "texte")]
    lignes = [[e.nom, e.prenom, e.classe, e.annee_scolaire, s.total_du + s.remise,
               s.remise, s.total_du, s.total_paye, s.solde, s.statut]
              for e, s in situations]
    total = ["TOTAL", None, None, None] + [
        sum(l[k] for l in lignes) for k in range(4, 9)] + [None]
    _ecrire(ws, colonnes, lignes, total)
    chemin = _chemin(dossier, "eleves")
    wb.save(chemin)
    return chemin


def exporter_impayes(resultat, dossier) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Impayés {resultat['date_ref']}"
    colonnes = [("Élève", "texte"), ("Classe", "texte"), ("Échéance impayée", "texte"),
                ("Date limite", "date"), ("Jours de retard", "nombre"),
                ("Exigible (FCFA)", "montant"), ("Payé (FCFA)", "montant"),
                ("En retard (FCFA)", "montant"), ("Parent / tuteur", "texte"),
                ("Téléphone", "texte")]
    lignes = [[l["nom"], l["classe"], l["tranche"], _date(l["date_limite"]), l["jours"],
               l["exigible"], l["paye"], l["retard"], l["tuteur"], l["telephone"]]
              for l in resultat["lignes"]]
    total = ["TOTAL", None, None, None, None] + [
        sum(l[k] for l in lignes) for k in range(5, 8)] + [None, None]
    _ecrire(ws, colonnes, lignes, total)
    chemin = _chemin(dossier, f"impayes_{resultat['date_ref']}")
    wb.save(chemin)
    return chemin


def exporter_cloture(cloture, dossier) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Par caissier"
    colonnes = ([("Caissier", "texte"), ("Paiements", "nombre")]
                + [(f"{libelle} (FCFA)", "montant") for libelle in MODES.values()]
                + [("Total (FCFA)", "montant")])
    lignes = [[c["nom"], c["nb"]] + [c["par_mode"][m] for m in MODES] + [c["total"]]
              for c in cloture["caissiers"]]
    total = (["TOTAL", cloture["nb"]] + [cloture["par_mode"][m] for m in MODES]
             + [cloture["total"]])
    _ecrire(ws, colonnes, lignes, total)

    detail = wb.create_sheet("Détail")
    colonnes_d = [("N° de reçu", "texte"), ("Date", "date"), ("Élève", "texte"),
                  ("Classe", "texte"), ("Mode", "texte"), ("Montant (FCFA)", "montant"),
                  ("Caissier", "texte"), ("État", "texte"),
                  ("Motif d'annulation", "texte")]
    paiements = ([p for c in cloture["caissiers"] for p in c["paiements"]]
                 + list(cloture["annules"]["paiements"]))
    paiements.sort(key=lambda p: (p["date_paiement"], p["numero_recu"] or ""))
    lignes_d = [[p["numero_recu"], _date(p["date_paiement"]), f"{p['nom']} {p['prenom']}",
                 p["classe"], MODES.get(p["mode"], p["mode"]), p["montant"],
                 p["encaisse_par"] or "Non renseigné",
                 "Annulé" if p["annule"] else "Valide", p["motif_annulation"] or ""]
                for p in paiements]
    _ecrire(detail, colonnes_d, lignes_d)
    chemin = _chemin(dossier, f"cloture_{cloture['debut']}_{cloture['fin']}")
    wb.save(chemin)
    return chemin