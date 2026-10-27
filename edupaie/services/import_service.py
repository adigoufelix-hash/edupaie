import re
import sqlite3
import unicodedata
from datetime import date
from itertools import islice
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from edupaie.services.echeancier_service import normaliser_telephone
from edupaie.services.paiement_service import ErreurMetier

MAX_LIGNES = 5000
MAX_ENTETE = 10  # l'en-tête doit se trouver dans les 10 premières lignes

ALIAS = {
    "nom": {"nom", "nom de famille"},
    "prenom": {"prenom", "prenoms"},
    "classe": {"classe", "niveau"},
    "annee": {"annee", "annee scolaire", "annee academique"},
    "total": {"total", "total du", "total du fcfa", "total des frais",
              "total des frais fcfa", "frais", "frais fcfa", "montant", "ecolage"},
    "tuteur": {"tuteur", "parent", "parent tuteur", "parent ou tuteur", "responsable"},
    "telephone": {"telephone", "tel", "telephone parent", "contact", "numero", "whatsapp"},
}
MESSAGE_LECTURE = ("Impossible de lire ce fichier. Vérifie qu'il s'agit bien d'un "
                   "classeur Excel (.xlsx), non protégé par un mot de passe.")


def _cle(texte) -> str:
    """Forme comparable : sans accents, sans majuscules, sans ponctuation."""
    texte = unicodedata.normalize("NFD", str(texte or "").casefold())
    texte = "".join(c for c in texte if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texte).split())


def _texte(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    return " ".join(str(valeur).split())


def _annee_courante() -> str:
    t = date.today()
    debut = t.year if t.month >= 8 else t.year - 1
    return f"{debut}-{debut + 1}"


def _annee(valeur, defaut: str) -> str:
    texte = _texte(valeur)
    if not texte:
        return defaut
    m = re.fullmatch(r"(\d{4})\s*[-/–]\s*(\d{4})", texte)
    if m:
        if int(m.group(2)) != int(m.group(1)) + 1:
            raise ValueError("Année scolaire incohérente (exemple : 2026-2027).")
        return f"{m.group(1)}-{m.group(2)}"
    if re.fullmatch(r"\d{4}", texte):
        return f"{texte}-{int(texte) + 1}"
    raise ValueError("Année scolaire invalide (exemple : 2026-2027).")


def _total(valeur):
    """Entier > 0, ou None si la cellule est vide."""
    if valeur is None or _texte(valeur) == "":
        return None
    if isinstance(valeur, bool):
        raise ValueError("Le total doit être un nombre entier.")
    if isinstance(valeur, (int, float)):
        if isinstance(valeur, float) and not valeur.is_integer():
            raise ValueError("Le total doit être un nombre entier.")
        nombre = int(valeur)
    else:
        brut = re.sub(r"[\s\u00a0\u202f]", "", str(valeur))
        brut = re.sub(r"(?i)fcfa|cfa|f$", "", brut)
        if not re.fullmatch(r"[0-9]+", brut):
            raise ValueError("Le total doit être un nombre entier, sans virgule ni point.")
        nombre = int(brut)
    if nombre <= 0:
        raise ValueError("Le total doit être supérieur à 0.")
    return nombre


class ImportEleveService:
    def __init__(self, repo, eleve_repo, frais_service=None, auth=None, sauvegarde=None):
        self.repo = repo
        self.eleves = eleve_repo
        self.frais = frais_service  # pour le total quand la colonne est vide
        self.auth = auth            # None dans les tests : pas de contrôle de droits
        self.sauvegarde = sauvegarde  # copie de sécurité faite avant l'import
        self.max_lignes = MAX_LIGNES

    def _exiger_droit(self):
        if self.auth is not None and not self.auth.peut("eleve.importer"):
            raise ErreurMetier("Seul un administrateur peut importer des élèves.")

    # --- Lecture du fichier --------------------------------------------------------
    def _lire(self, chemin):
        chemin = Path(chemin)
        if chemin.suffix.lower() != ".xlsx":
            raise ErreurMetier("Format non pris en charge : enregistre ton fichier "
                               "au format Excel (.xlsx).")
        try:
            classeur = load_workbook(chemin, read_only=True, data_only=True)
        except Exception:
            raise ErreurMetier(MESSAGE_LECTURE) from None
        try:
            return list(islice(classeur.worksheets[0].iter_rows(min_row=1, values_only=True),
                               self.max_lignes + MAX_ENTETE + 2))
        except Exception:
            raise ErreurMetier(MESSAGE_LECTURE) from None
        finally:
            classeur.close()

    @staticmethod
    def _entete(lignes):
        for i, ligne in enumerate(lignes[:MAX_ENTETE]):
            colonnes = {}
            for j, cellule in enumerate(ligne):
                cle = _cle(cellule)
                for champ, alias in ALIAS.items():
                    if cle in alias and champ not in colonnes:
                        colonnes[champ] = j
            if {"nom", "prenom", "classe"} <= colonnes.keys():
                return i, colonnes
        raise ErreurMetier(
            "En-têtes introuvables : le fichier doit contenir au minimum les colonnes "
            "Nom, Prénom et Classe. Télécharge le modèle Excel pour partir d'un "
            "fichier correct.")

    # --- Validation ----------------------------------------------------------------
    @staticmethod
    def _cle_eleve(nom, prenom, classe, annee):
        return (_cle(nom), _cle(prenom), _cle(classe), annee)

    def _cles_existantes(self):
        return {self._cle_eleve(e.nom, e.prenom, e.classe, e.annee_scolaire)
                for e in self.eleves.lister("", None)}

    def _valider(self, brut, defaut_annee):
        nom, prenom, classe = (_texte(brut.get(c)) for c in ("nom", "prenom", "classe"))
        if not nom:
            raise ValueError("Nom manquant.")
        if not prenom:
            raise ValueError("Prénom manquant.")
        if not classe:
            raise ValueError("Classe manquante.")
        if len(nom) > 60 or len(prenom) > 60:
            raise ValueError("Nom ou prénom trop long (60 caractères au maximum).")
        if len(classe) > 30:
            raise ValueError("Nom de classe trop long (30 caractères au maximum).")
        annee = _annee(brut.get("annee"), defaut_annee)
        total = _total(brut.get("total"))
        if total is None:
            total = self.frais.total(classe, annee) if self.frais is not None else 0
            if not total:
                raise ValueError("Total non renseigné et aucun frais défini pour cette "
                                 "classe : renseigne le total ou définis les frais.")
        tuteur = _texte(brut.get("tuteur"))
        if len(tuteur) > 80:
            raise ValueError("Nom du parent trop long (80 caractères au maximum).")
        telephone = _texte(brut.get("telephone"))
        if telephone and normaliser_telephone(telephone) is None:
            raise ValueError("Téléphone invalide (8 chiffres pour le Togo, ou avec "
                             "l'indicatif).")
        return {"nom": nom, "prenom": prenom, "classe": classe, "annee": annee,
                "total": total, "tuteur": tuteur, "telephone": telephone}

    def analyser(self, chemin) -> dict:
        """Lit et vérifie le fichier sans rien enregistrer."""
        self._exiger_droit()
        lignes = self._lire(chemin)
        entete, colonnes = self._entete(lignes)
        corps = lignes[entete + 1:]
        if len(corps) > self.max_lignes:
            raise ErreurMetier(f"Le fichier dépasse {self.max_lignes} lignes : "
                               "découpe-le en plusieurs fichiers.")
        defaut_annee = _annee_courante()
        existants, vus, resultat = self._cles_existantes(), set(), []
        for numero, ligne in enumerate(corps, start=entete + 2):  # numéro de ligne Excel
            if not any(_texte(cellule) for cellule in ligne):
                continue  # ligne vide
            brut = {champ: ligne[col] if col < len(ligne) else None
                    for champ, col in colonnes.items()}
            entree = {"ligne": numero, "nom": _texte(brut.get("nom")),
                      "prenom": _texte(brut.get("prenom")),
                      "classe": _texte(brut.get("classe")),
                      "annee": _texte(brut.get("annee")),
                      "total": _texte(brut.get("total")),
                      "tuteur": _texte(brut.get("tuteur")),
                      "telephone": _texte(brut.get("telephone")),
                      "statut": "ok", "message": ""}
            try:
                entree.update(self._valider(brut, defaut_annee))
            except ValueError as e:
                entree.update(statut="erreur", message=str(e))
            else:
                cle = self._cle_eleve(entree["nom"], entree["prenom"],
                                      entree["classe"], entree["annee"])
                if cle in existants:
                    entree.update(statut="ignore", message="Déjà enregistré : ignoré.")
                elif cle in vus:
                    entree.update(statut="ignore", message="Doublon dans le fichier : ignoré.")
                else:
                    vus.add(cle)
            resultat.append(entree)
        return {"fichier": Path(chemin).name, "lignes": resultat,
                "nb_ok": sum(l["statut"] == "ok" for l in resultat),
                "nb_ignores": sum(l["statut"] == "ignore" for l in resultat),
                "nb_erreurs": sum(l["statut"] == "erreur" for l in resultat)}

    # --- Import --------------------------------------------------------------------
    def importer(self, analyse) -> int:
        """Enregistre les lignes valides, en une seule transaction."""
        self._exiger_droit()
        existants, vus, a_importer = self._cles_existantes(), set(), []
        for l in analyse["lignes"]:
            if l["statut"] != "ok":
                continue
            cle = self._cle_eleve(l["nom"], l["prenom"], l["classe"], l["annee"])
            if cle in existants or cle in vus:  # la base a pu changer depuis l'aperçu
                continue
            vus.add(cle)
            a_importer.append(l)
        if not a_importer:
            raise ErreurMetier("Aucun élève à importer.")
        if self.sauvegarde is not None:
            try:
                self.sauvegarde.sauvegarder("avant-import")
            except (ErreurMetier, OSError, sqlite3.Error) as e:
                raise ErreurMetier("Import annulé : la sauvegarde de sécurité a "
                                   f"échoué ({e}).") from None
        try:
            return self.repo.ajouter_lot(a_importer)
        except sqlite3.Error as e:
            raise ErreurMetier("Import annulé, aucune donnée n'a été modifiée : "
                               f"{e}") from None

    # --- Modèle Excel --------------------------------------------------------------
    def creer_modele(self, dossier) -> Path:
        self._exiger_droit()
        classeur = Workbook()
        feuille = classeur.active
        feuille.title = "Élèves"
        entetes = ["Nom", "Prénom", "Classe", "Année scolaire", "Total dû (FCFA)",
                   "Parent / tuteur", "Téléphone"]
        for j, titre in enumerate(entetes, start=1):
            c = feuille.cell(row=1, column=j, value=titre)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E1B4B")
            c.alignment = Alignment(horizontal="center")
            feuille.column_dimensions[chr(64 + j)].width = 22
        feuille.freeze_panes = "A2"
        aide = classeur.create_sheet("Mode d'emploi")
        lignes = [
            "Mode d'emploi de l'import des élèves",
            "",
            "Remplis la feuille « Élèves » : une ligne par élève, sous les en-têtes (ne les modifie pas).",
            "Obligatoire : Nom, Prénom, Classe.",
            "Année scolaire : par exemple 2026-2027. Si la colonne est vide, l'année en cours est utilisée.",
            "Total dû (FCFA) : un nombre entier. S'il est vide, on utilise les frais définis pour la classe.",
            "Parent / tuteur et Téléphone : facultatifs. Téléphone : 8 chiffres (Togo) ou avec l'indicatif.",
            "Un élève déjà enregistré (même nom, prénom, classe et année) est ignoré : tu peux réimporter un fichier corrigé.",
            "Les lignes en erreur sont signalées avant l'import et ne sont pas importées.",
            "Une sauvegarde de la base est faite automatiquement juste avant l'import.",
            "Enregistre le fichier au format Excel (.xlsx).",
            "",
            "Exemple de ligne : KOUASSI | Kofi | 6e A | 2026-2027 | 100000 | M. Kouassi | 90 12 34 56",
        ]
        for i, texte in enumerate(lignes, start=1):
            aide.cell(row=i, column=1, value=texte)
        aide["A1"].font = Font(bold=True, size=14)
        aide.column_dimensions["A"].width = 110
        dossier = Path(dossier)
        dossier.mkdir(parents=True, exist_ok=True)
        chemin = dossier / "modele_import_eleves.xlsx"
        classeur.save(chemin)
        return chemin