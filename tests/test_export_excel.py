import tempfile
import unittest
from datetime import date

from openpyxl import load_workbook

from edupaie.models.entities import Eleve
from edupaie.services.export_excel import (exporter_cloture, exporter_eleves,
                                           exporter_impayes)
from edupaie.services.paiement_service import Situation


class TestExportExcel(unittest.TestCase):
    def test_export_des_eleves(self):
        situations = [
            (Eleve(None, "=1+1", "Un", "6e A", "2026-2027", 100000),
             Situation(90000, 40000, 50000, "Partiellement payé", 10000)),
            (Eleve(None, "B", "Deux", "6e B", "2026-2027", 50000),
             Situation(50000, 50000, 0, "Soldé", 0)),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            ws = load_workbook(exporter_eleves(situations, tmp)).active
        self.assertEqual(ws["A1"].value, "Nom")
        self.assertEqual(ws["A2"].value, "=1+1")  # reste un texte, jamais une formule
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual((ws["E2"].value, ws["G2"].value, ws["I2"].value),
                         (100000, 90000, 50000))
        self.assertEqual((ws["A4"].value, ws["E4"].value, ws["H4"].value),
                         ("TOTAL", 150000, 90000))

    def test_export_des_impayes(self):
        resultat = {"date_ref": "2026-10-25", "sans_echeancier": 0, "total_retard": 40000,
                    "lignes": [{"nom": "Kouassi Kofi", "classe": "6e A",
                                "tranche": "1re tranche", "date_limite": "2026-10-15",
                                "jours": 10, "exigible": 40000, "paye": 0, "retard": 40000,
                                "tuteur": "M. Kouassi", "telephone": "+228 90 12 34 56"}]}
        with tempfile.TemporaryDirectory() as tmp:
            ws = load_workbook(exporter_impayes(resultat, tmp)).active
        self.assertEqual(ws.title, "Impayés 2026-10-25")
        self.assertEqual(ws["D2"].value.date(), date(2026, 10, 15))
        self.assertEqual((ws["H2"].value, ws["H3"].value), (40000, 40000))
        self.assertEqual(ws["J2"].value, "+228 90 12 34 56")

    def test_export_de_la_cloture(self):
        paiement = {"numero_recu": "REC-2026-000001", "date_paiement": "2026-09-01",
                    "mode": "especes", "montant": 40000, "encaisse_par": "Caissier",
                    "annule": 0, "motif_annulation": None, "nom": "Kouassi",
                    "prenom": "Kofi", "classe": "6e A"}
        annule = dict(paiement, numero_recu="REC-2026-000002", montant=10000,
                      annule=1, motif_annulation="Montant saisi par erreur")
        cloture = {"debut": "2026-09-01", "fin": "2026-09-30", "nb": 1, "total": 40000,
                   "par_mode": {"especes": 40000, "cheque": 0, "virement": 0,
                                "mobile_money": 0},
                   "caissiers": [{"nom": "Caissier", "nb": 1, "total": 40000,
                                  "par_mode": {"especes": 40000, "cheque": 0,
                                               "virement": 0, "mobile_money": 0},
                                  "paiements": [paiement]}],
                   "annules": {"nb": 1, "total": 10000, "paiements": [annule]}}
        with tempfile.TemporaryDirectory() as tmp:
            wb = load_workbook(exporter_cloture(cloture, tmp))
        self.assertEqual(wb.sheetnames, ["Par caissier", "Détail"])
        self.assertEqual(wb["Par caissier"]["G3"].value, 40000)
        detail = wb["Détail"]
        self.assertEqual(detail.max_row, 3)
        self.assertEqual((detail["H2"].value, detail["H3"].value), ("Valide", "Annulé"))
        self.assertEqual(detail["I3"].value, "Montant saisi par erreur")


if __name__ == "__main__":
    unittest.main()